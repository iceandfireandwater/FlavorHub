from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path
import re

src=Path(r'E:\Desktop\agent面试题图片\FlavorHub_技术八股专项.md')
out=Path(r'E:\Desktop\agent面试题图片\FlavorHub_技术八股专项.docx')
lines=src.read_text(encoding='utf-8').splitlines()
doc=Document(); sec=doc.sections[0]
sec.top_margin=Inches(.8); sec.bottom_margin=Inches(.75); sec.left_margin=Inches(.85); sec.right_margin=Inches(.85)
sec.header_distance=Inches(.35); sec.footer_distance=Inches(.35)
styles=doc.styles
n=styles['Normal']; n.font.name='Calibri'; n._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); n.font.size=Pt(10.5); n.paragraph_format.space_after=Pt(5); n.paragraph_format.line_spacing=1.12
for name,size,color,before,after in [('Heading 1',16,'1F4E79',14,7),('Heading 2',13,'2F6690',11,5),('Heading 3',11.5,'365F91',8,4)]:
 s=styles[name]; s.font.name='Calibri'; s._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor.from_string(color); s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after); s.paragraph_format.keep_with_next=True
lead=styles.add_style('Lead',WD_STYLE_TYPE.PARAGRAPH); lead.base_style=n; lead.font.size=Pt(10.5); lead.paragraph_format.space_after=Pt(5)

def shade(cell,fill):
 tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)
def margins(cell):
 tcPr=cell._tc.get_or_add_tcPr(); tcMar=OxmlElement('w:tcMar')
 for k,v in [('top',120),('start',160),('bottom',120),('end',160)]:
  e=OxmlElement('w:'+k); e.set(qn('w:w'),str(v)); e.set(qn('w:type'),'dxa'); tcMar.append(e)
 tcPr.append(tcMar)
def add_text(p,s):
 s=s.replace('**',''); s=re.sub(r'`([^`]*)`',r'\1',s); p.add_run(s)

hp=sec.header.paragraphs[0]; hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT; rr=hp.add_run('FlavorHub · 技术八股专项'); rr.font.size=Pt(8.5); rr.font.color.rgb=RGBColor(107,114,128)
fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER; rr=fp.add_run('FlavorHub · 面试准备'); rr.font.size=Pt(8.5); rr.font.color.rgb=RGBColor(107,114,128)
for text,size,color,bold,italic,after in [('FlavorHub',24,'1F4E79',True,False,2),('技术八股专项面试题库',16,'365F91',True,False,3),('基于 Multi-Agent 与 LangGraph 的智能问答系统',10.5,'6B7280',False,True,12)]:
 p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(after); r=p.add_run(text); r.font.size=Pt(size); r.font.color.rgb=RGBColor.from_string(color); r.bold=bold; r.italic=italic

i=0
while i<len(lines):
 line=lines[i]
 if not line.strip() or line.strip()=='---' or line.startswith('# FlavorHub'): i+=1; continue
 if line.startswith('> '):
  t=doc.add_table(rows=1,cols=1); c=t.cell(0,0); shade(c,'EEF4F8'); margins(c); p=c.paragraphs[0]; p.style='Lead'; add_text(p,line[2:]); i+=1; continue
 if line.startswith('## '): p=doc.add_paragraph(style='Heading 1'); add_text(p,line[3:]); i+=1; continue
 if line.startswith('### '): p=doc.add_paragraph(style='Heading 2'); add_text(p,line[4:]); i+=1; continue
 if line.startswith('#### '): p=doc.add_paragraph(style='Heading 3'); add_text(p,line[5:]); i+=1; continue
 if line.startswith('- '): p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.left_indent=Inches(.28); p.paragraph_format.first_line_indent=Inches(-.18); add_text(p,line[2:]); i+=1; continue
 if re.match(r'^\d+\.\s',line): p=doc.add_paragraph(style='List Number'); p.paragraph_format.left_indent=Inches(.28); p.paragraph_format.first_line_indent=Inches(-.18); add_text(p,re.sub(r'^\d+\.\s','',line)); i+=1; continue
 if line.startswith('```') or line.startswith('~~~'):
  fence=line[:3]; i+=1; code=[]
  while i<len(lines) and not lines[i].startswith(fence): code.append(lines[i]); i+=1
  if i<len(lines): i+=1
  p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.18); p.paragraph_format.right_indent=Inches(.18); p.paragraph_format.space_after=Pt(6); pPr=p._p.get_or_add_pPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),'F3F4F6'); pPr.append(shd); r=p.add_run('\n'.join(code)); r.font.name='Consolas'; r.font.size=Pt(9); continue
 p=doc.add_paragraph(); add_text(p,line.strip()); i+=1

doc.core_properties.title='FlavorHub 技术八股专项面试题库'; doc.core_properties.subject='Multi-Agent、LangGraph、混合检索与评测'; doc.core_properties.author='Codex'; doc.save(out); print(out)